#!/usr/bin/env python3
"""Globalizace — premium Apple-Keynote-style Czech deck (10 slides)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- palette ----------
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
MIST    = RGBColor(0xD6, 0xDA, 0xE0)   # soft white for body
DIM     = RGBColor(0xA1, 0xA8, 0xB3)   # dimmed labels
ACCENT  = RGBColor(0x2E, 0x9B, 0xFF)   # Apple blue
GREEN   = RGBColor(0x32, 0xD7, 0x4B)   # positive
AMBER   = RGBColor(0xFF, 0x9F, 0x0A)   # caution
RED     = RGBColor(0xFF, 0x45, 0x3A)   # negative
INK     = RGBColor(0x0A, 0x0A, 0x0C)

HEAD = "Helvetica Neue"   # Apple feel; substitutes cleanly on Windows
BODY = "Helvetica Neue"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

# ---------- helpers ----------
def add_bg(slide, path):
    pic = slide.shapes.add_picture(path, 0, 0, SW, SH)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)
    return pic

def _set_tracking(run, val):
    """letter-spacing in 1/100 pt."""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(val)))

def _no_autofit(tf):
    tf.word_wrap = True

def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

def style_run(r, text, size, color, font=BODY, bold=False, tracking=None,
              italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    r.font.bold = bold
    r.font.italic = italic
    if tracking is not None:
        _set_tracking(r, tracking)

def para(tf, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    return p

def kicker(slide, text, color=ACCENT, x=Inches(0.95), y=Inches(0.85), w=Inches(8)):
    tb, tf = textbox(slide, x, y, w, Inches(0.5))
    p = para(tf, True)
    style_run(p.add_run(), text.upper(), 14.5, color, HEAD, bold=True, tracking=340)
    return tb

def title(slide, text, x=Inches(0.92), y=Inches(1.4), w=Inches(8.6),
          size=50, color=WHITE):
    tb, tf = textbox(slide, x, y, w, Inches(2.0))
    p = para(tf, True)
    p.line_spacing = 1.0
    style_run(p.add_run(), text, size, color, HEAD, bold=True, tracking=-30)
    return tb

def accent_rule(slide, x, y, w=Inches(0.62), color=ACCENT, h=Pt(3.2)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, int(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def points(slide, items, x=Inches(0.95), y=Inches(3.5), w=Inches(7.4),
           gap=Inches(1.02), dot=ACCENT, head_size=21, sub_size=14.5):
    """Each item: (heading, subtitle)."""
    for i, (head, sub) in enumerate(items):
        cy = y + Emu(int(gap) * i)
        # marker dot
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, cy + Inches(0.10),
                                   Inches(0.13), Inches(0.13))
        d.fill.solid(); d.fill.fore_color.rgb = dot
        d.line.fill.background(); d.shadow.inherit = False
        tb, tf = textbox(slide, x + Inches(0.4), cy, w, gap)
        p1 = para(tf, True)
        style_run(p1.add_run(), head, head_size, WHITE, HEAD, bold=True, tracking=-10)
        if sub:
            p2 = para(tf)
            p2.space_before = Pt(2)
            style_run(p2.add_run(), sub, sub_size, MIST, BODY, tracking=10)

def big_stat(slide, number, label, x=Inches(8.7), y=Inches(2.5),
             w=Inches(4.0), color=WHITE, num_size=110):
    tb, tf = textbox(slide, x, y, w, Inches(2.4))
    p = para(tf, True); p.alignment = PP_ALIGN.LEFT
    style_run(p.add_run(), number, num_size, color, HEAD, bold=True, tracking=-40)
    p2 = para(tf); p2.space_before = Pt(2)
    style_run(p2.add_run(), label.upper(), 14, DIM, HEAD, bold=True, tracking=260)
    return tb

def page_num(slide, n):
    tb, tf = textbox(slide, Inches(12.4), Inches(6.95), Inches(0.7), Inches(0.4))
    p = para(tf, True); p.alignment = PP_ALIGN.RIGHT
    style_run(p.add_run(), f"{n:02d}", 12, DIM, HEAD, bold=True, tracking=120)

def footer(slide, text="GLOBALIZACE"):
    tb, tf = textbox(slide, Inches(0.95), Inches(6.95), Inches(6), Inches(0.4))
    p = para(tf, True)
    style_run(p.add_run(), text, 11.5, DIM, HEAD, bold=True, tracking=300)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

BG = "assets/bg/%s.jpg"

# ============================================================ SLIDE 1 — TITLE
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s01_earth")
# top kicker centered
tb, tf = textbox(s, Inches(0), Inches(2.35), SW, Inches(0.5))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "PREZENTACE · 2026", 15, RGBColor(0xCF,0xE6,0xFF), HEAD,
          bold=True, tracking=420)
# giant title centered
tb, tf = textbox(s, Inches(0.5), Inches(2.9), Inches(12.33), Inches(1.7))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "Globalizace", 100, WHITE, HEAD, bold=True, tracking=-60)
# thin centered rule
rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.27), Inches(4.78),
                          Inches(0.8), Pt(3))
rule.fill.solid(); rule.fill.fore_color.rgb = ACCENT
rule.line.fill.background(); rule.shadow.inherit = False
# subtitle
tb, tf = textbox(s, Inches(0), Inches(5.0), SW, Inches(0.6))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "Tom Častvaj  ·  3.B", 22, MIST, BODY, tracking=60)
tb, tf = textbox(s, Inches(0), Inches(5.62), SW, Inches(0.5))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "Jak se svět propojil v jednu síť", 15.5, DIM, BODY,
          italic=True, tracking=20)
notes(s, "Úvod: Přivítejte třídu jednou silnou větou — 'Svět, ve kterém žijeme, "
         "je propojenější než kdykoli v historii.' Telefon v mé kapse, oblečení, "
         "jídlo i informace pocházejí z desítek zemí. Dnes si ukážeme, co tento jev "
         "znamená, proč vznikl a jaké má světlé i stinné stránky. Cíl: po prezentaci "
         "byste měli umět globalizaci vysvětlit kamarádovi jednou větou.")

# ============================================================ SLIDE 2 — CO JE
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s02_network")
kicker(s, "01 — Definice")
title(s, "Co je globalizace?")
accent_rule(s, Inches(0.95), Inches(2.62))
points(s, [
    ("Propojování světa v jeden celek",
     "Ekonomiky, kultury a lidé se stále těsněji prolínají."),
    ("Vzdálenosti přestaly platit",
     "Zboží i informace obletí planetu za hodiny, ne za měsíce."),
    ("Týká se úplně každého",
     "Od telefonu v kapse až po jídlo na talíři."),
], y=Inches(3.15), gap=Inches(1.12))
footer(s); page_num(s, 2)
notes(s, "Definice: Globalizace je proces, při kterém se státy, ekonomiky a kultury "
         "stále silněji propojují a stávají vzájemně závislými. Jednoduše: svět se "
         "'zmenšuje' — to, co se stane na jednom konci planety, dnes okamžitě "
         "ovlivní druhý. Proč to dnes řešíme? Protože tempo propojení je "
         "bezprecedentní: přes 5 miliard lidí je online a každý den se obchoduje "
         "napříč kontinenty. Příklad pro studenty: trend z TikToku se za den rozšíří "
         "do celého světa.")

# ============================================================ SLIDE 3 — JAK FUNGUJE
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s03_port")
kicker(s, "02 — Mechanismy")
title(s, "Jak to celé funguje")
accent_rule(s, Inches(0.95), Inches(2.62))
points(s, [
    ("Mezinárodní obchod", "Lodě a kontejnery propojují trhy všech kontinentů."),
    ("Internet", "Data a služby putují mezi zeměmi okamžitě."),
    ("Doprava", "Letecká i námořní logistika zkrátila vzdálenosti."),
    ("Globální komunikace", "Spolupracujeme v reálném čase přes časová pásma."),
], y=Inches(3.05), gap=Inches(0.95), head_size=19.5)
footer(s); page_num(s, 3)
notes(s, "Mechanismy globalizace — čtyři pilíře. 1) Mezinárodní obchod: kolem 80 % "
         "objemu světového zboží se přepraví po moři v kontejnerech. 2) Internet: "
         "umožňuje firmě v Praze spolupracovat s týmem v Indii během vteřin. "
         "3) Doprava: kontejnerová přeprava a letecká logistika radikálně zlevnily "
         "pohyb zboží. 4) Komunikace: videohovory a cloud propojují lidi napříč "
         "časovými pásmy. Tyto čtyři síly se navzájem posilují.")

# ============================================================ SLIDE 4 — PŘÍČINY
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s04_tech")
kicker(s, "03 — Příčiny")
title(s, "Co globalizaci pohání")
accent_rule(s, Inches(0.95), Inches(2.62))
points(s, [
    ("Technologický pokrok", "Počítače a chytré telefony spojily celé generace."),
    ("Rychlejší doprava", "Levnější a dostupnější přesun lidí i zboží."),
    ("Rozmach internetu", "Z lokální informace se během chvíle stane globální."),
    ("Mezinárodní byznys", "Firmy hledají trhy a partnery po celém světě."),
], y=Inches(3.05), gap=Inches(0.95), head_size=19.5)
footer(s); page_num(s, 4)
notes(s, "Příčiny: Proč se svět propojil právě teď? Klíčem je technologie. "
         "Vynález internetu (90. léta) a později chytrých telefonů (po roce 2007) "
         "umožnil okamžitou komunikaci pro miliardy lidí. Kontejnerová doprava "
         "zlevnila přepravu zboží na zlomek původní ceny. A firmy přirozeně hledají "
         "levnější výrobu a nové zákazníky za hranicemi. Souhra těchto faktorů "
         "globalizaci v posledních desetiletích extrémně urychlila.")

# ============================================================ SLIDE 5 — PŘÍKLADY (iPhone)
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s05_iphone")
kicker(s, "04 — V praxi", x=Inches(0.95))
title(s, "Globalizace v jednom\ntelefonu", size=44, w=Inches(7.2))
accent_rule(s, Inches(0.95), Inches(3.05))
points(s, [
    ("iPhone je dílem desítek zemí",
     "Design v USA, čipy z Asie, montáž v Číně."),
    ("Globální značky všude kolem",
     "Stejné logo potkáte v Tokiu i v Brně."),
    ("Mezinárodní dodavatelské řetězce",
     "Jeden výrobek, stovky propojených dodavatelů."),
], x=Inches(0.95), y=Inches(3.45), w=Inches(6.6), gap=Inches(0.95),
   head_size=18.5, sub_size=13.5)
big_stat(s, "40+", "zemí v jednom iPhonu", x=Inches(8.85), y=Inches(2.7),
         num_size=120, color=WHITE)
footer(s); page_num(s, 5)
notes(s, "Konkrétní příklad, který studenti znají — iPhone. Je navržen v Kalifornii, "
         "ale jeho součástky pocházejí od dodavatelů z více než 40 zemí: displeje a "
         "paměti z Jižní Koreje a Japonska, čipy z Tchaj-wanu, finální montáž v Číně. "
         "Žádná jediná země by ho sama nevyrobila tak levně a kvalitně. Stejně fungují "
         "globální značky (McDonald's, Nike, Zara) a celé dodavatelské řetězce. "
         "Otázka do třídy: Kolik zemí asi 'leží' ve vašem tričku?")

# ============================================================ SLIDE 6 — VÝHODY
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s06_market")
kicker(s, "05 — Výhody", color=GREEN)
title(s, "Co nám přináší")
accent_rule(s, Inches(0.95), Inches(2.62), color=GREEN)
points(s, [
    ("Větší výběr zboží", "Produkty z celého světa na dosah ruky."),
    ("Přístup k informacím", "Vědění planety dostupné komukoli online."),
    ("Ekonomický růst", "Obchod vytváří pracovní místa a bohatství."),
    ("Mezinárodní spolupráce", "Společné řešení problémů přesahujících hranice."),
], y=Inches(3.05), gap=Inches(0.95), head_size=19.5, dot=GREEN)
footer(s); page_num(s, 6)
notes(s, "Výhody — pozitivní strana. 1) Výběr: v běžném obchodě najdeme ovoce z Jižní "
         "Ameriky i elektroniku z Asie. 2) Informace: díky internetu má student v "
         "Česku přístup ke stejným kurzům jako student v New Yorku. 3) Růst: "
         "mezinárodní obchod (přes 24 bilionů dolarů ročně) pomohl vytáhnout stovky "
         "milionů lidí z chudoby. 4) Spolupráce: globální výzvy — pandemie, klima — "
         "lze řešit jen společně. Zdůrazněte: globalizace není jen ekonomika, je to i "
         "sdílení hodnot.")

# ============================================================ SLIDE 7 — DALŠÍ VÝHODY
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s07_science")
# right-aligned content
kicker(s, "06 — Výhody II", color=GREEN, x=Inches(6.6), w=Inches(6))
title(s, "Motor pokroku", x=Inches(6.55), w=Inches(6.2))
accent_rule(s, Inches(6.6), Inches(2.62), color=GREEN)
points(s, [
    ("Inovace", "Konkurence i spolupráce ženou nové nápady vpřed."),
    ("Vědecký pokrok", "Objevy se okamžitě sdílejí napříč kontinenty."),
    ("Vzdělání a znalosti", "Kurzy a data otevřené komukoli, kdekoli."),
], x=Inches(6.6), y=Inches(3.2), w=Inches(6.0), gap=Inches(1.05),
   head_size=20, dot=GREEN)
footer(s); page_num(s, 7)
notes(s, "Další výhody — dlouhodobý dopad. Inovace: globální konkurence nutí firmy "
         "neustále vylepšovat produkty (srovnejte telefon dnes a před 15 lety). "
         "Vědecký pokrok: vývoj vakcín proti covidu byl dílem mezinárodních týmů a "
         "trval měsíce místo let — protože vědci sdíleli data online. Vzdělání: "
         "platformy jako Khan Academy nebo Wikipedie zpřístupnily znalosti zdarma "
         "komukoli s internetem. Globalizace tak zrychluje lidské poznání jako celek.")

# ============================================================ SLIDE 8 — NEVÝHODY
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s08_local")
kicker(s, "07 — Nevýhody", color=AMBER)
title(s, "Druhá strana mince")
accent_rule(s, Inches(0.95), Inches(2.62), color=AMBER)
points(s, [
    ("Tlak na místní firmy", "Malé obchody těžko soupeří s globálními giganty."),
    ("Závislost na trzích", "Krize v cizině zasáhne ekonomiku i u nás doma."),
    ("Přesun pracovních míst", "Výroba se stěhuje tam, kde je levnější."),
], y=Inches(3.15), gap=Inches(1.12), head_size=21, dot=AMBER)
footer(s); page_num(s, 8)
notes(s, "Nevýhody — stinná strana. 1) Místní firmy: malé rodinné obchody často "
         "nemohou cenově konkurovat nadnárodním řetězcům a e-shopům, a zanikají. "
         "2) Závislost: když se zasekne jediný dodavatelský řetězec (např. nedostatek "
         "čipů v roce 2021), pocítí to celý svět. 3) Práce: továrny se přesouvají do "
         "zemí s nižšími mzdami — to znamená ztrátu pracovních míst v původní zemi. "
         "Vyvažte: nejde o to globalizaci odsoudit, ale vidět její náklady.")

# ============================================================ SLIDE 9 — DALŠÍ NEVÝHODY
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s09_pollution")
kicker(s, "08 — Nevýhody II", color=RED)
title(s, "Skrytá cena propojení")
accent_rule(s, Inches(0.95), Inches(2.62), color=RED)
points(s, [
    ("Dopady na životní prostředí", "Doprava zboží přes půl světa zvyšuje emise."),
    ("Stírání kulturních rozdílů", "Lokální tradice ustupují globálnímu mainstreamu."),
    ("Ekonomická nerovnost", "Zisky se nerozdělují mezi všechny rovnoměrně."),
], y=Inches(3.15), gap=Inches(1.12), head_size=21, dot=RED)
footer(s); page_num(s, 9)
notes(s, "Další nevýhody — méně viditelné, zato zásadní. 1) Životní prostředí: "
         "převoz zboží přes celou planetu znamená obrovskou spotřebu paliv; námořní "
         "doprava sama stojí za významnou částí globálních emisí CO₂. 2) Kultura: "
         "stejné značky, filmy a móda všude vedou ke 'kulturní uniformitě' — mizí "
         "místní svébytnost. 3) Nerovnost: z globalizace netěží všichni stejně — "
         "rozdíly mezi bohatými a chudými mohou růst. Pointa: propojení má cenu, "
         "kterou není vždy vidět na cenovce.")

# ============================================================ SLIDE 10 — ZÁVĚR
s = prs.slides.add_slide(BLANK)
add_bg(s, BG % "s10_sunrise")
tb, tf = textbox(s, Inches(0), Inches(1.5), SW, Inches(0.5))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "09 — ZÁVĚR", 14.5, RGBColor(0xFF,0xE0,0xC0), HEAD,
          bold=True, tracking=360)
tb, tf = textbox(s, Inches(0.5), Inches(2.05), Inches(12.33), Inches(1.2))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "Svět bez hranic?", 58, WHITE, HEAD, bold=True, tracking=-40)
# balanced summary
tb, tf = textbox(s, Inches(2.0), Inches(3.45), Inches(9.33), Inches(1.0))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.18
style_run(p.add_run(),
          "Globalizace nás spojuje, obohacuje a zrychluje pokrok — "
          "zároveň přináší závislost, nerovnost a tlak na přírodu.",
          18.5, MIST, BODY, tracking=10)
# key takeaway pill
pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.67), Inches(4.78),
                          Inches(6.0), Inches(0.62))
pill.adjustments[0] = 0.5
pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
pill.fill.fore_color.rgb = ACCENT
pill.line.fill.background(); pill.shadow.inherit = False
ptf = pill.text_frame; ptf.word_wrap = True
pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
style_run(pp.add_run(), "Nejde ji zastavit — jde ji dělat chytře.", 16, WHITE,
          HEAD, bold=True, tracking=20)
# discussion question
tb, tf = textbox(s, Inches(1.5), Inches(5.75), Inches(10.33), Inches(1.0))
p = para(tf, True); p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), "Otázka pro vás:  ", 16.5, DIM, HEAD, bold=True, tracking=40)
style_run(p.add_run(),
          "Je globalizace spíš příležitost, nebo hrozba — a pro koho?",
          16.5, WHITE, BODY, italic=True, tracking=10)
notes(s, "Závěr — vyvážené shrnutí. Globalizace není ani jednoznačně dobrá, ani "
         "špatná — je to mocný nástroj se dvěma tvářemi. Spojuje nás, dává nám výběr "
         "a žene pokrok; zároveň vytváří závislosti, prohlubuje nerovnost a zatěžuje "
         "planetu. Klíčové poselství: globalizaci nelze zastavit, ale můžeme ji "
         "směrovat zodpovědně (férový obchod, udržitelná doprava, ochrana kultur). "
         "Zakončete otevřenou otázkou a vyzvěte spolužáky k diskusi — nechte je "
         "hlasovat: příležitost, nebo hrozba? Děkuji za pozornost.")

prs.save("Globalizace.pptx")
print("saved Globalizace.pptx —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
