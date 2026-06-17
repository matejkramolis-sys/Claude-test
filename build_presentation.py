#!/usr/bin/env python3
"""Builds 'The Vessel, New York City' PowerPoint.
Clean Apple-style design: white background, subtle gray accents,
large fonts, image-dominant layouts, speaker notes on every slide.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import copy

# ---- palette ----
INK   = RGBColor(0x1D, 0x1D, 0x1F)   # near-black Apple text
GRAY  = RGBColor(0x6E, 0x6E, 0x73)   # secondary text
LGRAY = RGBColor(0x86, 0x86, 0x8B)   # tertiary / kicker
HAIR  = RGBColor(0xD2, 0xD2, 0xD7)   # subtle divider
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEARW = RGBColor(0xF5, 0xF5, 0xF7)   # Apple light gray panel
CREDIT= RGBColor(0xB0, 0xB0, 0xB5)

FONT  = "Helvetica Neue"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    # white background
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def cover(s, path, x, y, w, h):
    """Add image cropped to fully cover the x,y,w,h box (no distortion)."""
    iw, ih = Image.open(path).size
    box_ar = w / h
    img_ar = iw / ih
    pic = s.shapes.add_picture(path, x, y, w, h)
    if img_ar > box_ar:                     # too wide -> crop sides
        crop = (1 - box_ar / img_ar) / 2
        pic.crop_left = crop; pic.crop_right = crop
    else:                                   # too tall -> crop top/bottom
        crop = (1 - img_ar / box_ar) / 2
        pic.crop_top = crop; pic.crop_bottom = crop
    pic.shadow.inherit = False
    return pic


def scrim(s, x, y, w, h, color=RGBColor(0, 0, 0), alpha=35):
    """Semi-transparent overlay rectangle for text legibility over photos."""
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    # set transparency
    sp = r.fill.fore_color._xFill.find(qn('a:srgbClr'))
    a = sp.makeelement(qn('a:alpha'), {'val': str((100 - alpha) * 1000)})
    sp.append(a)
    return r


def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def para(tf, text, size, color, bold=False, first=False, align=PP_ALIGN.LEFT,
         space_after=8, space_before=0, line=None, tracking=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    r = p.add_run(); r.text = text
    f = r.font
    f.name = FONT; f.size = Pt(size); f.bold = bold
    f.color.rgb = color
    if tracking is not None:
        rPr = r._r.get_or_add_rPr(); rPr.set('spc', str(tracking))
    return p


def credit(s, text, light=False):
    tf = tb(s, Inches(0.5), Inches(7.12), Inches(12.33), Inches(0.3))
    color = RGBColor(0xEC, 0xEC, 0xEE) if light else CREDIT
    para(tf, text, 8.5, color, first=True)


def kicker(tf, text, color=LGRAY, first=False):
    para(tf, text.upper(), 15, color, bold=True, first=first,
         space_after=10, tracking=300)


def badge(s, x, y, d, num):
    """Small gray circle number badge (clean icon style)."""
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    c.fill.solid(); c.fill.fore_color.rgb = NEARW
    c.line.color.rgb = HAIR; c.line.width = Pt(1)
    c.shadow.inherit = False
    tf = c.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num)
    r.font.name = FONT; r.font.size = Pt(20); r.font.bold = True
    r.font.color.rgb = INK


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# =====================================================================
# SLIDE 1 — TITLE  (full-bleed photo)
# =====================================================================
s = slide()
cover(s, "assets/title.jpg", 0, 0, SW, SH)
scrim(s, 0, 0, SW, SH, alpha=32)
scrim(s, 0, Inches(4.3), SW, Inches(3.2), alpha=28)
tf = tb(s, Inches(0.9), Inches(4.55), Inches(11.5), Inches(2.4))
para(tf, "THE VESSEL", 90, WHITE, bold=True, first=True, space_after=2, tracking=100)
para(tf, "New York City", 34, RGBColor(0xEA, 0xEA, 0xEC), space_after=0)
credit(s, "Photo: Gunnar Klack / Wikimedia Commons, CC BY-SA 4.0", light=True)
notes(s, "Hello everyone, and welcome. Today I want to show you one of the most "
        "unusual buildings in New York City. It is called The Vessel. As you can "
        "see in this big photo, it does not look like a normal building. It has no "
        "rooms and no offices. It is made only of stairs that you can climb. It "
        "looks a little bit like a giant honeycomb or a basket made of copper. "
        "In the next few minutes, I will tell you where it is, who built it, why "
        "it is so special, and why it became famous all over the world. Let us start.")

# =====================================================================
# SLIDE 2 — WHERE IS IT?  (text + photo + map)
# =====================================================================
s = slide()
LX, LW = Inches(0.9), Inches(4.4)
tf = tb(s, LX, Inches(1.0), LW, Inches(5.5))
kicker(tf, "Location", first=True)
para(tf, "Where Is It?", 46, INK, bold=True, space_after=22, line=1.0)
para(tf, "Hudson Yards", 26, INK, bold=True, space_after=4)
para(tf, "West side of Manhattan, New York City", 21, GRAY, space_after=18, line=1.1)
para(tf, "Next to the Hudson River, close to the High Line park.",
     21, GRAY, space_after=0, line=1.2)
# big photo top-right, map bottom-right
RX = Inches(5.7); RW = SW - RX - Inches(0.5)
cover(s, "assets/context.jpg", RX, Inches(0.6), RW, Inches(3.9))
cover(s, "assets/map.png", RX, Inches(4.7), RW, Inches(2.3))
credit(s, "Photos: dconvertini / Wikimedia, CC BY-SA 2.0  ·  Map © OpenStreetMap contributors")
notes(s, "So, where can you find The Vessel? It stands in a part of New York called "
        "Hudson Yards. This is on the west side of Manhattan, very close to the "
        "Hudson River. You can see the city skyline in the top photo, and a map "
        "below it. The location is important. A few years ago, this area was old "
        "and mostly empty, with railway tracks. Then the city built a huge new "
        "neighbourhood with tall glass towers, shops, and restaurants. The Vessel "
        "was placed right in the middle, like a piece of art. It is also next to "
        "the High Line, a famous park built on an old train line, so many tourists "
        "walk past it every day.")

# =====================================================================
# SLIDE 3 — WHAT IS THE VESSEL?  (facts, photo right)
# =====================================================================
s = slide()
cover(s, "assets/facts.jpg", Inches(7.0), 0, SW - Inches(7.0), SH)
tf = tb(s, Inches(0.9), Inches(0.7), Inches(5.6), Inches(6.0))
kicker(tf, "Basic Facts", first=True)
para(tf, "What Is\nThe Vessel?", 42, INK, bold=True, space_after=16, line=0.98)
def fact(label, value):
    para(tf, label.upper(), 13, LGRAY, bold=True, space_after=1, tracking=200)
    para(tf, value, 23, INK, bold=True, space_after=13, line=1.0)
fact("Height", "About 46 metres (16 floors)")
fact("Opened", "March 2019")
fact("Steps", "154 stairs · 2,500 steps · 80 landings")
fact("Purpose", "Art you can climb — just for fun and views")
credit(s, "Photo: Brian W. Schaller / Wikimedia Commons, Free Art License")
notes(s, "Now, what exactly is The Vessel? The idea is simple but strange. It is a "
        "big art object that you can walk inside and climb. It is about 46 metres "
        "tall, which is like a 16-floor building. It opened in March 2019. Inside, "
        "there are 154 flights of stairs, around 2,500 steps, and 80 landings where "
        "you can stop and rest. The British designer Thomas Heatherwick made it. "
        "He wanted to build something that was not a normal tower with offices. "
        "Instead, the only purpose of The Vessel is to be enjoyed. People climb up, "
        "take photos, and look at the city and the river. So it is really art that "
        "you can use with your feet, not just look at.")

# =====================================================================
# SLIDE 4 — UNIQUE DESIGN  (text + two tall photos)
# =====================================================================
s = slide()
tf = tb(s, Inches(0.9), Inches(0.95), Inches(4.4), Inches(5.6))
kicker(tf, "Architecture", first=True)
para(tf, "A Unique\nDesign", 46, INK, bold=True, space_after=22, line=0.98)
para(tf, "Shaped like a giant honeycomb.", 22, INK, bold=True, space_after=14, line=1.15)
para(tf, "154 staircases connect together and climb in a circle.",
     21, GRAY, space_after=12, line=1.2)
para(tf, "Copper-coloured steel makes it shine in the sun.",
     21, GRAY, space_after=0, line=1.2)
IX = Inches(5.55); IW = (SW - IX - Inches(0.5) - Inches(0.2)) / 2
cover(s, "assets/design.jpg", IX, Inches(0.6), IW, Inches(6.3))
cover(s, "assets/design2.jpg", IX + IW + Inches(0.2), Inches(0.6), IW, Inches(6.3))
credit(s, "Photos: Mike Peel (www.mikepeel.net) / Wikimedia Commons, CC BY-SA 4.0")
notes(s, "Let us look closer at the design, because this is the most exciting part. "
        "Why does it look so different from normal buildings? Most buildings are "
        "boxes with walls and windows. The Vessel has no walls at all. It is open, "
        "like a basket. Its shape is wider at the top and smaller at the bottom. "
        "From far away it looks like a giant honeycomb or a beehive. The 154 "
        "staircases connect to each other, so you can walk up in many different "
        "ways and always find new views. The metal is painted a warm copper colour, "
        "so when the sun shines, the whole structure looks golden and bright. "
        "Because of this special shape, people often say it looks like a piece of "
        "modern art, not a building.")

# =====================================================================
# SLIDE 5 — FUN FACTS  (badges + photo)
# =====================================================================
s = slide()
cover(s, "assets/fun.jpg", Inches(8.1), 0, SW - Inches(8.1), SH)
tf = tb(s, Inches(0.9), Inches(0.85), Inches(6.6), Inches(1.4))
kicker(tf, "Did You Know?", first=True)
para(tf, "Fun Facts", 46, INK, bold=True, space_after=0)
facts = [
    "It cost about 200 million dollars to build.",
    "It weighs around 600 tonnes of steel.",
    "It was made in Italy, then shipped to New York.",
    "At first, climbing it was completely free.",
]
y = 2.35
for i, t in enumerate(facts, 1):
    badge(s, Inches(0.95), Inches(y), Inches(0.62), i)
    tf2 = tb(s, Inches(1.85), Inches(y - 0.05), Inches(5.5), Inches(0.9),
             anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, t, 22, INK, first=True, line=1.1, space_after=0)
    y += 1.12
credit(s, "Photo: Mike Peel (www.mikepeel.net) / Wikimedia Commons, CC BY-SA 4.0")
notes(s, "Here are some fun facts that surprise most people. First, The Vessel was "
        "very expensive. It cost about 200 million dollars. That is a lot of money "
        "for something with no rooms inside! Second, it is very heavy. It uses "
        "around 600 tonnes of steel. Third, the pieces were not built in New York. "
        "They were made in Italy and then sent by ship across the ocean, and "
        "workers put them together like a giant model kit. And fourth, when it "
        "first opened, you did not have to pay anything to climb it. Later they "
        "added a small ticket. These facts show how much work and money people "
        "put into one piece of public art.")

# =====================================================================
# SLIDE 6 — WHY PEOPLE VISIT  (full-bleed photo + scrim text)
# =====================================================================
s = slide()
cover(s, "assets/visit.jpg", 0, 0, SW, SH)
scrim(s, 0, 0, SW, SH, alpha=22)
scrim(s, 0, Inches(4.6), SW, Inches(2.9), alpha=30)
tf = tb(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(2.5))
kicker(tf, "A Tourist Favourite", color=RGBColor(0xE6, 0xE6, 0xE8), first=True)
para(tf, "Why People Visit", 48, WHITE, bold=True, space_after=10)
para(tf, "Amazing views  ·  Perfect photos  ·  A famous landmark",
     24, RGBColor(0xEA, 0xEA, 0xEC), space_after=0)
credit(s, "Photo: Kidfly182 / Wikimedia Commons, CC BY 4.0", light=True)
notes(s, "So why do so many people come here? There are three main reasons. The "
        "first reason is the views. When you climb to the top, you can see the tall "
        "buildings of New York and the Hudson River. The second reason is "
        "photography. The Vessel is full of repeating stairs and lines, so photos "
        "taken here look amazing. It quickly became one of the most photographed "
        "places in the city, and you can see it in thousands of pictures on social "
        "media. The third reason is simply that it is famous. People want to see it "
        "because everyone is talking about it. In its first years, millions of "
        "tourists came to climb it and take a selfie.")

# =====================================================================
# SLIDE 7 — PROBLEMS AND CONTROVERSIES  (photo + text)
# =====================================================================
s = slide()
cover(s, "assets/problems.jpg", Inches(6.9), 0, SW - Inches(6.9), SH)
tf = tb(s, Inches(0.9), Inches(1.0), Inches(5.5), Inches(5.6))
kicker(tf, "Controversies", first=True)
para(tf, "Problems", 46, INK, bold=True, space_after=22, line=1.0)
para(tf, "Sadly, some people were hurt here.", 24, INK, bold=True,
     space_after=14, line=1.15)
para(tf, "Because of safety worries, The Vessel was closed in 2021.",
     21, GRAY, space_after=12, line=1.2)
para(tf, "People said the walls and railings were too low and not safe enough.",
     21, GRAY, space_after=0, line=1.2)
credit(s, "Photo: Kidfly182 / Wikimedia Commons, CC BY 4.0")
notes(s, "Not everything about The Vessel is happy. It also had serious problems. "
        "Sadly, a number of people died here after falling from the high levels. "
        "Because of this, the owners decided to close The Vessel in 2021, so no "
        "one could go inside. Many people said the design was not safe enough. The "
        "side walls and railings were quite low, and that made the high floors "
        "dangerous. There was also a lot of discussion in the news. Some people "
        "asked if it was right to build something so beautiful but not safe. This "
        "was a difficult and sad time for the building. I will keep this part "
        "short, but it is important to understand the full story.")

# =====================================================================
# SLIDE 8 — THE VESSEL TODAY  (photo + text)
# =====================================================================
s = slide()
cover(s, "assets/today.jpg", 0, 0, Inches(7.0), SH)
tf = tb(s, Inches(7.5), Inches(1.0), Inches(5.3), Inches(5.6))
kicker(tf, "Right Now", first=True)
para(tf, "The Vessel\nToday", 46, INK, bold=True, space_after=22, line=0.98)
para(tf, "It reopened in October 2024.", 24, INK, bold=True,
     space_after=14, line=1.15)
para(tf, "New floor-to-ceiling steel nets now keep visitors safe.",
     21, GRAY, space_after=12, line=1.2)
para(tf, "The lower levels are open again, so people can climb and enjoy the views.",
     21, GRAY, space_after=0, line=1.2)
credit(s, "Photo: Kidfly182 / Wikimedia Commons, CC BY 4.0")
notes(s, "Now let us look at The Vessel today. After being closed for about three "
        "years, it reopened to the public in October 2024. But this time it is "
        "different. The owners added new safety features. They put strong steel "
        "nets, from the floor to the ceiling, on almost every level. These nets "
        "are hard to cut or climb, so visitors are much safer now. You can change "
        "a little of the view through the mesh, but most people think safety is "
        "more important. The lower levels are open again, and the very top floor "
        "stays closed. So today, people can once again walk up The Vessel, take "
        "photos, and enjoy the city — but in a safer way.")

# =====================================================================
# SLIDE 9 — MY OPINION  (photo + pros/cons)
# =====================================================================
s = slide()
cover(s, "assets/opinion.jpg", 0, 0, Inches(5.4), SH)
tf = tb(s, Inches(6.0), Inches(0.9), Inches(6.7), Inches(1.3))
kicker(tf, "My Opinion", first=True)
para(tf, "What I Think", 46, INK, bold=True, space_after=0)
# Pros column
px = Inches(6.0)
tfp = tb(s, px, Inches(2.5), Inches(3.2), Inches(4.0))
para(tfp, "+  GOOD", 18, INK, bold=True, first=True, space_after=12, tracking=200)
for t in ["Beautiful and very original", "Great views and photos", "Free public art for the city"]:
    para(tfp, t, 20, GRAY, space_after=12, line=1.15)
# Cons column
tfc = tb(s, Inches(9.55), Inches(2.5), Inches(3.2), Inches(4.0))
para(tfc, "–  NOT SO GOOD", 18, INK, bold=True, first=True, space_after=12, tracking=200)
for t in ["Safety was a big problem", "Very expensive to build", "The new nets change the view"]:
    para(tfc, t, 20, GRAY, space_after=12, line=1.15)
# divider
ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.4), Inches(2.55), Pt(1.2), Inches(3.4))
ln.fill.solid(); ln.fill.fore_color.rgb = HAIR; ln.line.fill.background(); ln.shadow.inherit = False
credit(s, "Photo: Dancerma / Wikimedia Commons, CC BY-SA 4.0")
notes(s, "Now I want to share my own opinion. For me, The Vessel is a building with "
        "two faces. On the good side, I think it is beautiful and very original. "
        "There is nothing else like it in the world, and the views and photos from "
        "the top are amazing. I also like that it started as free public art that "
        "everyone could enjoy. But on the other side, there are real problems. The "
        "safety issues were very serious and very sad, and we must not forget that. "
        "It was also extremely expensive, and now the new safety nets change the "
        "open feeling of the design. So, in my opinion, The Vessel is a brave and "
        "amazing idea, but it shows that beautiful design must always be safe first. "
        "What do you think?")

# =====================================================================
# SLIDE 10 — THANK YOU  (full-bleed photo)
# =====================================================================
s = slide()
cover(s, "assets/thanks.jpg", 0, 0, SW, SH)
scrim(s, 0, 0, SW, SH, alpha=42)
tf = tb(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2.2),
        anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Thank You", 80, WHITE, bold=True, first=True, align=PP_ALIGN.CENTER,
     space_after=6, tracking=100)
para(tf, "For Listening", 32, RGBColor(0xEA, 0xEA, 0xEC), align=PP_ALIGN.CENTER,
     space_after=0)
credit(s, "Photo: Kidfly182 / Wikimedia Commons, CC BY 4.0", light=True)
notes(s, "That is the end of my presentation. Thank you very much for listening. I "
        "hope you now know more about The Vessel in New York City — what it is, why "
        "it is special, and the problems it had. If you ever visit New York, go and "
        "see it for yourself. Do you have any questions? I am happy to answer them.")

prs.save("The_Vessel_NYC.pptx")
print("Saved The_Vessel_NYC.pptx with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
